"""One finished job = one applications.xlsx row, built from cv_data.json."""

import json
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from jobstitch.cv_renderer import TailoredCVData, WorkExperienceItem
from jobstitch.workbook import STATUS_VALUES, ApplicationsWorkbook

RESOURCES = Path(__file__).resolve().parents[1] / "resources"


def cv_data(company="Globex", job_title="Staff **Platform** Engineer"):
    return TailoredCVData(
        company_name=company,
        job_title=job_title,
        summary="A summary.",
        skills=[f"skill {i}" for i in range(6)],
        experiences=[
            WorkExperienceItem(
                title="Engineer",
                company="Initech",
                location="Rome, Italy",
                dates="2020 -- 2024",
                bullet_points=["Did a thing."],
            )
            for _ in range(3)
        ],
    )


def make_job(tmp_path, *, analysis=None, name="Globex_Staff_Engineer", **kwargs):
    folder = tmp_path / name
    folder.mkdir(parents=True)
    (folder / "cv_data.json").write_text(cv_data(**kwargs).model_dump_json())
    if analysis is not None:
        (folder / "analysis.json").write_text(json.dumps(analysis))
    return folder


def workbook(tmp_path):
    return ApplicationsWorkbook(
        xlsx_path=tmp_path / "applications.xlsx", resources_dir=RESOURCES
    )


def read_row(xlsx_path, row=2):
    ws = openpyxl.load_workbook(xlsx_path).active
    headers = [str(c.value).strip().lower() for c in ws[1]]
    return dict(zip(headers, (ws.cell(row=row, column=c + 1).value
                              for c in range(len(headers)))))


def test_seeds_the_spreadsheet_and_fills_it_from_cv_data(tmp_path):
    wb = workbook(tmp_path)
    assert not wb.xlsx_path.exists()

    wb.record(make_job(tmp_path))

    row = read_row(wb.xlsx_path)
    assert row["company_name"] == "Globex"
    # The highlighter's **bold** markers are LaTeX, not part of the title.
    assert row["job_title"] == "Staff Platform Engineer"
    assert row["application_date"] == date.today().isoformat()
    assert row["application status"] == STATUS_VALUES[0] == "not applied"
    assert row["notes"] is None
    assert row["weblink"] is None


def test_weblink_comes_from_analysis_json_when_a_source_left_one(tmp_path):
    wb = workbook(tmp_path)
    job = make_job(tmp_path, analysis={"posting_url": "https://jobs.example/1"})

    wb.record(job)

    assert read_row(wb.xlsx_path)["weblink"] == "https://jobs.example/1"


def test_an_unparsable_analysis_json_costs_the_link_and_nothing_else(tmp_path):
    wb = workbook(tmp_path)
    job = make_job(tmp_path)
    (job / "analysis.json").write_text("{not json")

    wb.record(job)

    row = read_row(wb.xlsx_path)
    assert row["weblink"] is None
    assert row["company_name"] == "Globex"


def test_rows_accumulate(tmp_path):
    wb = workbook(tmp_path)

    wb.record(make_job(tmp_path, name="first", company="Globex"))
    wb.record(make_job(tmp_path, name="second", company="Initech"))

    assert read_row(wb.xlsx_path, row=2)["company_name"] == "Globex"
    assert read_row(wb.xlsx_path, row=3)["company_name"] == "Initech"


def test_without_cv_data_there_is_no_row_to_write(tmp_path):
    """The caller decides what a failure costs; for the watcher it is a log
    line, so the spreadsheet must not be seeded on a job it cannot record."""
    wb = workbook(tmp_path)
    empty_job = tmp_path / "no_cv"
    empty_job.mkdir()

    with pytest.raises(FileNotFoundError):
        wb.record(empty_job)
    assert not wb.xlsx_path.exists()
